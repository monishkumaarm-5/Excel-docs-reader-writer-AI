# fileReader.py
"""Read Excel and Word files into the app's data structures."""

import os

from doxl_ai_terminal.data_structure.docs import DocData, DocLine
from doxl_ai_terminal.data_structure.excel import ExcelData, ExcelSheet, ExcelCell
from doxl_ai_terminal.data_structure.formatting import (
    base_format_for_paragraph,
    read_paragraph_alignment,
    read_paragraph_style,
    read_excel_cell_format,
)


def read_word(filepath) -> DocData:
    from docx import Document

    try:
        doc = Document(filepath)
    except Exception as e:
        raise ValueError(f"Could not open Word document '{filepath}': {e}") from e

    doc_data = DocData(
        filename=os.path.basename(filepath),
        total_paragraphs=0,
        total_lines=0,
    )

    for p_index, para in enumerate(doc.paragraphs, start=1):
        if para.text.strip() == "":
            continue

        fmt = base_format_for_paragraph(para)
        alignment = read_paragraph_alignment(para)
        para_style = read_paragraph_style(para)

        for l_index, line in enumerate(para.text.split("\n"), start=1):
            if line.strip() == "":
                continue
            doc_data.lines.append(DocLine(
                paragraph=p_index,
                line=l_index,
                line_str=line.strip(),
                font_name=fmt.get("font_name"),
                font_size=fmt.get("font_size"),
                bold=fmt.get("bold"),
                italic=fmt.get("italic"),
                underline=fmt.get("underline"),
                strikethrough=fmt.get("strikethrough"),
                font_color=fmt.get("font_color"),
                highlight_color=fmt.get("highlight_color"),
                superscript=fmt.get("superscript"),
                subscript=fmt.get("subscript"),
                alignment=alignment,
                style=para_style.get("style"),
                heading_level=para_style.get("heading_level"),
            ))
            doc_data.total_lines += 1

        doc_data.total_paragraphs += 1

    return doc_data


def read_excel(filepath) -> ExcelData:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    try:
        wb = load_workbook(filepath, read_only=True)
    except Exception as e:
        raise ValueError(f"Could not open Excel file '{filepath}': {e}") from e

    excel_data = ExcelData(
        filename=os.path.basename(filepath),
        total_sheets=len(wb.sheetnames),
    )

    for sheet in wb:
        sheet_data = ExcelSheet(
            sheet_name=sheet.title,
            total_rows=0,
            total_columns=0,
        )

        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    fmt = read_excel_cell_format(cell)
                    # row is stored as int (not str) — this is the key
                    # data structure fix. All internal code uses int rows.
                    sheet_data.cells.append(ExcelCell(
                        row=cell.row,                         # int
                        column=get_column_letter(cell.column), # str
                        data_excel=cell.value,
                        **fmt,
                    ))
                    sheet_data.total_rows = max(sheet_data.total_rows, cell.row)
                    sheet_data.total_columns = max(sheet_data.total_columns, cell.column)

        excel_data.sheets.append(sheet_data)

    wb.close()
    return excel_data
