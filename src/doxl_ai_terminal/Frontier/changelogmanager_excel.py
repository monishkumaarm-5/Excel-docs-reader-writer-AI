# changelogmanager_excel.py
"""Live Excel editing manager — read, edit, and sync an Excel file."""

from typing import List, Dict

from doxl_ai_terminal.Frontier.fileReader import read_excel
from doxl_ai_terminal.data_structure.change_log import ChangeLog, ChangeType
from doxl_ai_terminal.data_structure.excel import ExcelCell, coerce_excel_value, parse_row
from doxl_ai_terminal.data_structure.excel_write import write_excel
from doxl_ai_terminal.pipeline.terminal_ui import info


class LiveExcelManager:
    """Read, edit, and sync an Excel file in real-time."""

    def __init__(self, filepath: str):
        self.filepath = filepath
        self.excel_data = read_excel(filepath)
        self.changelog = ChangeLog()
        total = sum(len(s.cells) for s in self.excel_data.sheets)
        info(f"Loaded: {filepath} ({total} cells)")

    # ---- View ----

    def view(self, sheet_name: str = None) -> str:
        lines = []
        for sheet in self.excel_data.sheets:
            if sheet_name and sheet.sheet_name != sheet_name:
                continue
            lines.append(f"\nSheet: {sheet.sheet_name}")
            for cell in sheet.cells:
                lines.append(f"  [{cell.address}] = {cell.data_excel}")
        return "\n".join(lines)

    # ---- Search ----

    def search(self, keyword: str) -> List[Dict]:
        results = []
        for sheet in self.excel_data.sheets:
            for cell in sheet.cells:
                if keyword.lower() in str(cell.data_excel).lower():
                    results.append({
                        "sheet": sheet.sheet_name,
                        "cell": cell.address,
                        "value": cell.data_excel,
                    })
        return results

    # ---- Formatting kwargs accepted by update_cell/add_cell/format_cell ----
    _FORMAT_FIELDS = (
        "font_name", "font_size", "bold", "italic", "underline",
        "font_color", "fill_color", "alignment", "number_format",
    )

    @classmethod
    def _apply_format_kwargs(cls, cell: ExcelCell, format_kwargs: dict) -> dict:
        applied = {}
        for key, val in format_kwargs.items():
            if key in cls._FORMAT_FIELDS and val is not None:
                setattr(cell, key, val)
                applied[key] = val
        return applied

    # ---- Update Cell (with auto-sync) ----

    def update_cell(self, sheet_name: str, row: int, column: str, new_value, auto_save=True, **format_kwargs):
        """Update a cell's value. Optional formatting kwargs may be passed."""
        row = parse_row(row)
        for sheet in self.excel_data.sheets:
            if sheet.sheet_name == sheet_name:
                for cell in sheet.cells:
                    if cell.row == row and cell.column == column:
                        old = cell.data_excel
                        cell.data_excel = coerce_excel_value(new_value)
                        applied = self._apply_format_kwargs(cell, format_kwargs)

                        self.changelog.add(
                            ChangeType.UPDATE,
                            {"sheet": sheet_name, "row": row, "column": column},
                            old_value=old,
                            new_value=new_value,
                        )

                        if auto_save:
                            self._sync_to_file()

                        suffix = f" ({applied})" if applied else ""
                        return f"Updated {sheet_name}[{column}{row}]: '{old}' → '{new_value}'{suffix}"

        return f"Not found: {sheet_name}[{column}{row}]"

    # ---- Format Cell only (with auto-sync) ----

    def format_cell(self, sheet_name: str, row: int, column: str, auto_save=True, **format_kwargs):
        """Change a cell's formatting without touching its value."""
        row = parse_row(row)
        for sheet in self.excel_data.sheets:
            if sheet.sheet_name == sheet_name:
                for cell in sheet.cells:
                    if cell.row == row and cell.column == column:
                        applied = self._apply_format_kwargs(cell, format_kwargs)
                        if not applied:
                            return f"No formatting fields given for {sheet_name}[{column}{row}]"

                        self.changelog.add(
                            ChangeType.UPDATE,
                            {"sheet": sheet_name, "row": row, "column": column, "format": True},
                            new_value=str(applied),
                        )

                        if auto_save:
                            self._sync_to_file()

                        return f"Formatted {sheet_name}[{column}{row}]: {applied}"

        return f"Not found: {sheet_name}[{column}{row}]"

    # ---- Add Cell (with auto-sync) ----

    def add_cell(self, sheet_name: str, row: int, column: str, value, auto_save=True, **format_kwargs):
        row = parse_row(row)
        for sheet in self.excel_data.sheets:
            if sheet.sheet_name == sheet_name:
                for cell in sheet.cells:
                    if cell.row == row and cell.column == column:
                        old = cell.data_excel
                        cell.data_excel = coerce_excel_value(value)
                        self._apply_format_kwargs(cell, format_kwargs)

                        self.changelog.add(
                            ChangeType.UPDATE,
                            {"sheet": sheet_name, "row": row, "column": column},
                            old_value=old,
                            new_value=value,
                        )

                        if auto_save:
                            self._sync_to_file()

                        return (f"Updated (cell already had a value): "
                                f"{sheet_name}[{column}{row}] '{old}' -> '{value}'")

                new_cell = ExcelCell(row=row, column=column, data_excel=coerce_excel_value(value))
                self._apply_format_kwargs(new_cell, format_kwargs)
                sheet.cells.append(new_cell)

                self.changelog.add(
                    ChangeType.ADD,
                    {"sheet": sheet_name, "row": row, "column": column},
                    new_value=value,
                )

                if auto_save:
                    self._sync_to_file()

                return f"Added {sheet_name}[{column}{row}] = '{value}'"

        return f"Sheet not found: {sheet_name}"

    # ---- Add Row (with auto-sync) ----

    def add_row(self, sheet_name: str, row_data: Dict[str, str], insert_at: int = None, auto_save=True):
        """Add a full row. row_data = {"A": "val1", "B": "val2", ...}

        If `insert_at` is given (e.g. 3), the new row is inserted at
        that position and every existing row at or after it shifts down
        by one. Without it, the row is appended after the current last row.
        """
        for sheet in self.excel_data.sheets:
            if sheet.sheet_name == sheet_name:
                if insert_at is not None:
                    new_row_num = parse_row(insert_at)
                    for cell in sheet.cells:
                        if cell.row >= new_row_num:
                            cell.row = cell.row + 1
                else:
                    max_row = 0
                    for cell in sheet.cells:
                        max_row = max(max_row, cell.row)
                    new_row_num = max_row + 1

                for col, val in row_data.items():
                    sheet.cells.append(ExcelCell(
                        row=new_row_num, column=col,
                        data_excel=coerce_excel_value(val),
                    ))

                self.changelog.add(
                    ChangeType.ADD,
                    {"sheet": sheet_name, "row": new_row_num},
                    new_value=str(row_data),
                )

                if auto_save:
                    self._sync_to_file()

                return f"Added row {new_row_num} to {sheet_name}: {row_data}"

        return f"Sheet not found: {sheet_name}"

    # ---- Add Column (insert with shift, with auto-sync) ----

    def add_column(self, sheet_name: str, column: str, values: Dict[int, str] = None,
                    shift: bool = True, auto_save=True):
        """Insert a column at `column` (e.g. "C"). By default every
        existing column at or after it shifts one place to the right.

        `values` maps row number (int) to the value for the new column,
        e.g. {1: "Header", 2: "x"}.
        """
        from openpyxl.utils import column_index_from_string, get_column_letter

        for sheet in self.excel_data.sheets:
            if sheet.sheet_name != sheet_name:
                continue

            insert_idx = column_index_from_string(column)

            if shift:
                for cell in sorted(sheet.cells, key=lambda c: -column_index_from_string(c.column)):
                    idx = column_index_from_string(cell.column)
                    if idx >= insert_idx:
                        cell.column = get_column_letter(idx + 1)

            values = values or {}
            for row_num, val in values.items():
                sheet.cells.append(ExcelCell(
                    row=parse_row(row_num), column=column,
                    data_excel=coerce_excel_value(val),
                ))

            self.changelog.add(
                ChangeType.ADD,
                {"sheet": sheet_name, "column": column},
                new_value=str(values),
            )

            if auto_save:
                self._sync_to_file()

            note = f" with {len(values)} value(s)" if values else " (empty)"
            return f"Added column {column} to {sheet_name}{note}"

        return f"Sheet not found: {sheet_name}"

    # ---- Delete Column (with auto-sync) ----

    def delete_column(self, sheet_name: str, column: str, shift: bool = True, auto_save=True):
        """Delete every cell in `column`."""
        from openpyxl.utils import column_index_from_string, get_column_letter

        for sheet in self.excel_data.sheets:
            if sheet.sheet_name != sheet_name:
                continue

            col_idx = column_index_from_string(column)
            to_remove = [c for c in sheet.cells if c.column == column]

            for cell in to_remove:
                sheet.cells.remove(cell)

            if shift:
                for cell in sorted(sheet.cells, key=lambda c: column_index_from_string(c.column)):
                    idx = column_index_from_string(cell.column)
                    if idx > col_idx:
                        cell.column = get_column_letter(idx - 1)

            self.changelog.add(
                ChangeType.DELETE,
                {"sheet": sheet_name, "column": column},
                old_value=f"{len(to_remove)} cells",
            )

            if auto_save:
                self._sync_to_file()

            return f"Deleted column {column} ({len(to_remove)} cells)"

        return f"Sheet not found: {sheet_name}"

    # ---- Delete Cell (with auto-sync) ----

    def delete_cell(self, sheet_name: str, row: int, column: str, auto_save=True):
        row = parse_row(row)
        for sheet in self.excel_data.sheets:
            if sheet.sheet_name == sheet_name:
                for cell in sheet.cells:
                    if cell.row == row and cell.column == column:
                        old = cell.data_excel
                        sheet.cells.remove(cell)

                        self.changelog.add(
                            ChangeType.DELETE,
                            {"sheet": sheet_name, "row": row, "column": column},
                            old_value=old,
                        )

                        if auto_save:
                            self._sync_to_file()

                        return f"Deleted {sheet_name}[{column}{row}]: '{old}'"

        return f"Not found: {sheet_name}[{column}{row}]"

    # ---- Delete Row (with auto-sync) ----

    def delete_row(self, sheet_name: str, row: int, auto_save=True):
        """Delete all cells in a row."""
        row = parse_row(row)
        for sheet in self.excel_data.sheets:
            if sheet.sheet_name == sheet_name:
                to_remove = [c for c in sheet.cells if c.row == row]

                if not to_remove:
                    return f"Row {row} is empty"

                for cell in to_remove:
                    sheet.cells.remove(cell)

                self.changelog.add(
                    ChangeType.DELETE,
                    {"sheet": sheet_name, "row": row},
                    old_value=f"{len(to_remove)} cells",
                )

                if auto_save:
                    self._sync_to_file()

                return f"Deleted row {row} ({len(to_remove)} cells)"

        return f"Sheet not found: {sheet_name}"

    # ---- Replace All (with auto-sync) ----

    def replace_all(self, old_text: str, new_text: str, auto_save=True):
        count = 0
        for sheet in self.excel_data.sheets:
            for cell in sheet.cells:
                cell_text = str(cell.data_excel)
                if old_text in cell_text:
                    old = cell.data_excel
                    cell.data_excel = cell_text.replace(old_text, new_text)
                    count += 1

                    self.changelog.add(
                        ChangeType.UPDATE,
                        {"sheet": sheet.sheet_name, "row": cell.row, "column": cell.column},
                        old_value=old,
                        new_value=cell.data_excel,
                    )

        if auto_save and count > 0:
            self._sync_to_file()

        return f"Replaced '{old_text}' → '{new_text}' in {count} cells"

    # ---- Batch Edit ----

    def batch_edit(self, edits: List[Dict], auto_save: bool = True):
        """Apply multiple edits and save once.

        Each edit: {"action": "update/add/delete", "sheet": str, "row": int, "column": str, "value": str}
        """
        results = []

        for edit in edits:
            action = edit["action"]
            sheet = edit["sheet"]
            row = parse_row(edit["row"])

            if action == "update":
                r = self.update_cell(sheet, row, edit["column"], edit["value"], auto_save=False)
            elif action == "add":
                r = self.add_cell(sheet, row, edit["column"], edit["value"], auto_save=False)
            elif action == "delete":
                r = self.delete_cell(sheet, row, edit["column"], auto_save=False)
            else:
                r = f"Unknown action: {action}"

            results.append(r)

        if auto_save:
            self._sync_to_file()
        return results

    # ---- Sync / Save / Reload ----

    def _sync_to_file(self):
        write_excel(self.excel_data, self.filepath)

    def save_as(self, output_path: str):
        write_excel(self.excel_data, output_path, source_path=self.filepath)
        return f"Saved copy: {output_path}"

    def reload(self):
        self.excel_data = read_excel(self.filepath)
        self.changelog.clear()
        return f"Reloaded from: {self.filepath}"

    def history(self) -> str:
        if not self.changelog.has_changes():
            return "No changes yet."
        return self.changelog.summary()
