#docs_write.py
import os
from typing import Optional

from doxl_ai_terminal.data_structure.docs import DocData
from doxl_ai_terminal.data_structure.formatting import (
    alignment_to_wd,
    apply_run_format,
    apply_paragraph_style,
    base_format_for_paragraph,
)


def write_word(doc_data: DocData, output_path: str, source_path: Optional[str] = None):
    """Write DocData back to a .docx file.

    This used to build a brand-new `docx.Document()` from scratch on every
    save, which threw away everything the data model doesn't track:
    original styles, headers/footers, tables, images, and any blank
    paragraph the reader had skipped. It also joined a paragraph's lines
    with a plain "\\n" and handed the whole thing to `add_paragraph`,
    losing per-line formatting even though the embedded "\\n" itself is
    turned into a real line break by python-docx.

    Instead, this loads the existing document (`source_path` if given,
    else `output_path` if it already exists) and edits paragraphs in
    place by index — `paragraph` numbers from read_word are exactly the
    1-based position in `doc.paragraphs`, blank paragraphs included, so
    this maps straight back onto the original document. Only paragraphs
    that actually hold tracked lines are touched; each line becomes its
    own run (separated by explicit `run.add_break()` calls) so lines can
    carry independent formatting, and `paragraph.clear()` preserves the
    paragraph's own style/alignment instead of discarding it.
    """
    from docx import Document

    load_path = source_path or (output_path if os.path.exists(output_path) else None)
    doc = Document(load_path) if load_path and os.path.exists(load_path) else Document()

    # Group lines by paragraph, keeping each line's L number alongside its
    # text (and formatting) so we can order by L rather than by insertion
    # order. Every tool/prompt tells the agent that L is a position
    # ("insert this as line 2"), so the saved file has to actually honor
    # that instead of just appending whatever was added most recently.
    para_map: dict[int, list] = {}
    for line in doc_data.lines:
        para_map.setdefault(line.paragraph, []).append(line)

    for p_num in sorted(para_map.keys()):
        ordered = sorted(para_map[p_num], key=lambda dl: dl.line)

        if 1 <= p_num <= len(doc.paragraphs):
            paragraph = doc.paragraphs[p_num - 1]
            base_format = base_format_for_paragraph(paragraph)
            paragraph.clear()  # drops runs only — style/alignment survive
        else:
            # A brand-new paragraph beyond what the source document had.
            # Pad with empty paragraphs so indices stay aligned with
            # read_word's numbering if there's a gap.
            while len(doc.paragraphs) < p_num - 1:
                doc.add_paragraph("")
            paragraph = doc.add_paragraph()
            base_format = {}

        requested_alignment = next((dl.alignment for dl in ordered if dl.alignment), None)
        if requested_alignment:
            mapped = alignment_to_wd(requested_alignment)
            if mapped is not None:
                paragraph.alignment = mapped

        # A paragraph becomes/stops being a heading only through style --
        # heading_level itself is derived, never written directly (see
        # formatting.apply_paragraph_style's docstring).
        requested_style = next((dl.style for dl in ordered if dl.style), None)
        if requested_style:
            apply_paragraph_style(paragraph, doc, requested_style)

        prev_run = None
        for dl in ordered:
            if prev_run is not None:
                prev_run.add_break()
            run = paragraph.add_run(dl.line_str)
            apply_run_format(run, dl, base_format)
            prev_run = run

    doc.save(output_path)
    print(f"Saved: {output_path}")
